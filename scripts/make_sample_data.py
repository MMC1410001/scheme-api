"""Write made-up files into local_data/ so the whole flow can run without GCS.

    uv run python scripts/make_sample_data.py
    uv run python scripts/make_sample_data.py --perf-dealers 20000   # extra data for k6 load tests

Scheme files mimic the SAP export (one CSV per day, only that day's updated rows).
Day 2 updates dealer 6100000001's LUXURY row, so API 1 should show the day-2 numbers.
Reports (API 2) mimic the generator's grouped workbook; v2 should be served, not v1.
"""
import argparse
import csv
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import openpyxl.utils

ROOT = Path(__file__).resolve().parent.parent / "local_data"
today = datetime.now(ZoneInfo("Asia/Kolkata")).date()
first = today.replace(day=1)
last = (first + timedelta(days=32)).replace(day=1) - timedelta(days=1)

HEADER = [
    "Client", "Agreement", "Customer", "Valid from", "To", "Deal Suffix", "Comp desc",
    "GSR (Q)", "Invoiced Quantity(Q)", "GSR (B)", "Rebate Rate", "Clear Cut CN",
    "Short Fall", "Short Fall %", "Last Updated Date", "Last Updated Time",
]


def d(day: date) -> str:
    return day.strftime("%d/%m/%Y")


def row(customer, suffix, comp, to, gsr_q, qty, gsr_b, rate, clear_cut, short, short_pct, updated):
    return [
        "100", f"02{customer[-8:]}", customer, d(first), d(to), suffix, comp, gsr_q, qty, gsr_b,
        rate, clear_cut, short, short_pct, d(updated), "10:20:56",
    ]


slab_end = {1: first.replace(day=11), 2: first.replace(day=18), 3: first.replace(day=25)}
day1, day2 = first, first + timedelta(days=1)
DAY1 = []
for customer, gsr in (("6100000001", "405232.86 "), ("6100000002", "882.00-")):  # SAP writes negatives as "882.00-"
    DAY1 += [
        row(customer, "03", "EB SLAB 1", slab_end[1], gsr, "0.000", gsr, "0.75", "0.00", "0.00", "0.00", day1),
        row(customer, "04", "EB SLAB 2", slab_end[2], gsr, "0.000", gsr, "0.75", "0.00", "0.00", "0.00", day1),
        row(customer, "05", "EB SLAB 3", slab_end[3], gsr, "0.000", gsr, "0.75", "0.00", "0.00", "0.00", day1),
        row(customer, "09", "VALUE", last, gsr, "0.000", gsr, "0.00", "0.00", "0.00", "0.00", day1),
        row(customer, "06", "LUXURY", last, "100.00", "200.000", "90000.00", "1.00", "900.00", "200.00", "50.00", day1),
        row(customer, "07", "PREMIUM 01", last, "0.00", "0.000", "0.00", "0.00", "0.00", "40.00", "100.00", day1),
        row(customer, "08", "PREMIUM 02", last, "0.00", "0.000", "0.00", "0.00", "0.00", "8.00", "100.00", day1),
        row(customer, "5!", "EB FULL MONTH", last, "0.00", "0.000", gsr, "2.25", "0.00", "0.00", "0.00", day1),
        row(customer, "65", "ALWD EB2", last, gsr, "0.000", gsr, "1.00", "0.00", "0.00", "0.00", day1),  # dropped
    ]
DAY2 = [
    row("6100000001", "06", "LUXURY", last, "315.10", "315.100", "159292.80", "1.00", "1592.93", "84.90", "21.23", day2),
]


def write_csv(rows, day):
    folder = ROOT / f"sap-exports/{today.year}/{today.month:02d}"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"sap-exports_{today.year}_{today.month:02d}_main_{day:%Y%m%d}_030003_001.csv"
    with path.open("w", newline="") as f:
        csv.writer(f).writerows([HEADER, *rows])
    print("wrote", path.relative_to(ROOT.parent))


def write_perf_files(count: int) -> None:
    """Dealers 6200000001.. with all 8 components, split into 50,000-row files like SAP does."""
    rows = []
    for n in range(1, count + 1):
        customer = f"62{n:08d}"
        gsr = f"{(n * 37) % 900000 + 100:.2f}"
        rows += [
            row(customer, "03", "EB SLAB 1", slab_end[1], gsr, "0.000", gsr, "0.75", "0.00", "0.00", "0.00", day1),
            row(customer, "04", "EB SLAB 2", slab_end[2], gsr, "0.000", gsr, "0.75", "0.00", "0.00", "0.00", day1),
            row(customer, "05", "EB SLAB 3", slab_end[3], gsr, "0.000", gsr, "0.75", "0.00", "0.00", "0.00", day1),
            row(customer, "09", "VALUE", last, gsr, "0.000", gsr, "0.00", "0.00", "0.00", "0.00", day1),
            row(customer, "06", "LUXURY", last, "100.00", "200.000", "90000.00", "1.00", "900.00", "20.00", "9.09", day1),
            row(customer, "07", "PREMIUM 01", last, "0.00", "0.000", "0.00", "0.00", "0.00", "40.00", "100.00", day1),
            row(customer, "08", "PREMIUM 02", last, "0.00", "0.000", "0.00", "0.00", "0.00", "8.00", "100.00", day1),
            row(customer, "5!", "EB FULL MONTH", last, "0.00", "0.000", gsr, "2.25", "0.00", "0.00", "0.00", day1),
        ]
    folder = ROOT / f"sap-exports/{today.year}/{today.month:02d}"
    folder.mkdir(parents=True, exist_ok=True)
    for part, start in enumerate(range(0, len(rows), 50_000), start=1):
        path = folder / f"sap-exports_{today.year}_{today.month:02d}_perf_{day1:%Y%m%d}_000000_{part:03d}.csv"
        with path.open("w", newline="") as f:
            csv.writer(f).writerows([HEADER, *rows[start:start + 50_000]])
        print("wrote", path.relative_to(ROOT.parent))


parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
parser.add_argument("--perf-dealers", type=int, default=0, help="also write N made-up dealers for load tests")
args = parser.parse_args()

write_csv(DAY1, day1)
write_csv(DAY2, day2)
if args.perf_dealers:
    write_perf_files(args.perf_dealers)

# API 2: reports in the generator's layout (made-up values): a group row merged over each
# group's columns, a column-name row, one row per dealer. v1 and v2 are written so the
# API can be checked to serve the highest version (v2 has GSR(Q) Value 125000.5 for 6100000001).
REPORT_GROUPS = [
    ("Value Slab", ["GSR(Q) Value", "GSR(B) Value", "Value CN%", "Value Clear Cut", "Short Fall Value",
                    "Short Fall % Value"]),
    ("Early Bird", ["GSR(Q) EB 1", "GSR(B) EB 1", "EB 1 CN%", "Short Fall EB 1", "Clear Cut"]),
    ("Luxury", ["Volume Q Luxury", "Luxury CN%", "Clear Cut"]),
]
REPORT_FIXED = ["Zone", "Region", "Depot Code", "Customer", "Customer Name"]


def report_row(customer, name, gsr_q):
    # 0.4000000059604645 and 66.86000061035156 are how the generator's 32-bit floats arrive;
    # the loader turns them back into 0.4 and 66.86. None is an empty cell.
    return [
        "Central", "Bihar North", "1302", customer, name,
        gsr_q, gsr_q, 0.4000000059604645, 0, 33431.359375, 66.86000061035156,
        336479.875, 336479.875, 0.75, 0, None,
        315, 1.0, 1592.93,
    ]


def write_report(version: int, rows) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Output"
    groups, names, col = list(REPORT_FIXED), [None] * len(REPORT_FIXED), len(REPORT_FIXED) + 1
    merges = [f"{openpyxl.utils.get_column_letter(c)}1:{openpyxl.utils.get_column_letter(c)}2"
              for c in range(1, col)]
    for group, columns in REPORT_GROUPS:
        groups += [group] + [None] * (len(columns) - 1)
        names += columns
        merges.append(f"{openpyxl.utils.get_column_letter(col)}1:"
                      f"{openpyxl.utils.get_column_letter(col + len(columns) - 1)}1")
        col += len(columns)
    sheet.append(groups)
    sheet.append(names)
    for cells in rows:
        sheet.append(cells)
    for cells in merges:
        sheet.merge_cells(cells)
    month_name = today.strftime("%B")
    folder = ROOT / f"automated-generated-files/output/{today.year}/{today.month:02d}"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / (f"automated-generated-files_output_{today.year}_{today.month:02d}_output_"
                     f"{month_name}_{today.year}_grouped_v{version}.xlsx")
    book.save(path)
    print("wrote", path.relative_to(ROOT.parent))


perf_report_rows = [report_row(f"62{n:08d}", f"Perf Dealer {n}", float((n * 37) % 900000 + 100))
                    for n in range(1, args.perf_dealers + 1)]
write_report(1, [report_row("6100000001", "Sample Paints", 120000.0),
                 report_row("6100000002", "Om Traders", -882.0)])
write_report(2, [report_row("6100000001", "Sample Paints", 125000.5),
                 report_row("6100000002", "Om Traders", -882.0), *perf_report_rows])
