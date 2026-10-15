import io

import openpyxl

from ingestion.excel import flatten_header, read_rows, read_table


def make_xlsx(build) -> bytes:
    workbook = openpyxl.Workbook()
    build(workbook.active)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_two_row_header_with_merged_cells_is_flattened():
    def build(sheet):
        sheet.append(["Dealer Code", "Value", None, "Premium", None])
        sheet.append([None, "GSR Q Value", "Benefit", "Invoice Qty", "Benefit"])
        sheet.append(["0001001", 1200, 50, 30, 10])
        sheet.merge_cells("A1:A2")  # vertical merge
        sheet.merge_cells("B1:C1")  # "Value" spans two columns
        sheet.merge_cells("D1:E1")

    rows = read_rows(make_xlsx(build), "schemes.xlsx", header_rows=2)

    assert rows == [
        {
            "dealer_code": "0001001",
            "value_gsr_q_value": 1200,
            "value_benefit": 50,
            "premium_invoice_qty": 30,
            "premium_benefit": 10,
        }
    ]


def test_blank_rows_are_skipped():
    def build(sheet):
        sheet.append(["Customer", "Scheme"])
        sheet.append([None, None])
        sheet.append(["1001", "Luxury"])

    assert read_rows(make_xlsx(build), "f.xlsx", header_rows=1) == [{"customer": "1001", "scheme": "Luxury"}]


def test_header_names_are_snake_case_unique_and_never_empty():
    assert flatten_header([["Short Fall %", "Benefit", "Benefit", None]]) == [
        "short_fall_pct",
        "benefit",
        "benefit_2",
        "column_4",
    ]


def test_csv_files_are_supported():
    content = "Customer,Scheme\n1001,EB Slab 1\n".encode()
    assert read_rows(content, "f.csv", header_rows=1) == [{"customer": "1001", "scheme": "EB Slab 1"}]


def test_group_names_fill_right_over_merged_cells():
    # Fixed columns merged over both header rows, then groups whose
    # name is only in the first cell. Repeated column names stay unique thanks to the group prefix.
    def build(sheet):
        sheet.append(["Zone", "Customer", "Value Slab", None, "Early Bird", None, None])
        sheet.append([None, None, "GSR(Q) Value", "Clear Cut", "GSR(Q) EB 1", "EB 1 CN%", "Clear Cut"])
        sheet.append(["Central", "6100000001", 1200, 0, 300, 0.75, None])
        for cells in ("A1:A2", "B1:B2", "C1:D1", "E1:G1"):
            sheet.merge_cells(cells)

    columns, rows = read_table(make_xlsx(build), "r.xlsx", header_rows=2)

    assert columns == [
        "zone", "customer", "value_slab_gsr_q_value", "value_slab_clear_cut",
        "early_bird_gsr_q_eb_1", "early_bird_eb_1_cn_pct", "early_bird_clear_cut",
    ]
    assert list(rows) == [["Central", "6100000001", 1200, 0, 300, 0.75, None]]


def test_short_rows_are_padded_to_the_header_width():
    content = "Customer,Scheme,Benefit\n1001,Luxury\n".encode()
    columns, rows = read_table(content, "f.csv", header_rows=1)
    assert list(rows) == [["1001", "Luxury", None]]
